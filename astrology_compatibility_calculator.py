import pandas as pd
import numpy as np
from astroquery.jplhorizons import Horizons
from astropy.time import Time
import warnings
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

warnings.filterwarnings("ignore", category=UserWarning)

# ===============================
# KONFIGURATION
# ===============================
GEBURTSDATUM = '1979-06-04 12:00'
JAHR = 1987

PLANET_IDS = {
    'Sonne': '10', 'Mond': '301', 'Merkur': '199', 'Venus': '299',
    'Mars': '499', 'Jupiter': '599', 'Saturn': '699'
}

PRIORITY_PLANETS = ['Sonne', 'Venus', 'Mars']

ASPECTS = [(0, 8, 100), (60, 6, 70), (90, 8, 25), (120, 8, 90), (180, 8, 40)]

ZODIAC_START = {
    "Wassermann": "01-20", "Fische": "02-19", "Widder": "03-21",
    "Stier": "04-20", "Zwillinge": "05-21", "Krebs": "06-21",
    "Löwe": "07-23", "Jungfrau": "08-23", "Waage": "09-23",
    "Skorpion": "10-23", "Schütze": "11-22", "Steinbock": "12-22"
}

# ===============================
# FUNKTIONEN (unverändert)
# ===============================
def fetch_birth_positions(date_str):
    jd = Time(date_str).jd
    pos = {}
    for name, pid in PLANET_IDS.items():
        eph = Horizons(id=pid, location='399', epochs=jd).ephemerides(quantities='31')
        pos[name] = float(eph['ObsEclLon'][0])
    return pos

def fetch_year_positions(year):
    epochs = {'start': f"{year}-01-01", 'stop': f"{year}-12-31", 'step': '1d'}
    base = Horizons(id='10', location='399', epochs=epochs).ephemerides()
    dates = pd.to_datetime(base['datetime_str'])
    data = {}
    for name, pid in PLANET_IDS.items():
        eph = Horizons(id=pid, location='399', epochs=epochs).ephemerides(quantities='31')
        data[name] = eph['ObsEclLon'].filled(np.nan).astype(float)
    return dates, data

def calculate_score(birth, transit):
    score_sum = 0
    hits = 0
    for b_name, b_lon in birth.items():
        for t_name, t_lon in transit.items():
            diff = abs(b_lon - t_lon) % 360
            if diff > 180: diff = 360 - diff
            for ang, orb, val in ASPECTS:
                if abs(diff - ang) <= orb:
                    weight = 1.0 + (0.25 if b_name in PRIORITY_PLANETS else 0) + (0.25 if t_name in PRIORITY_PLANETS else 0)
                    score_sum += val * weight
                    hits += 1
                    break
    return round((score_sum / hits) * np.log1p(hits), 2) if hits > 0 else 50.0

def get_zodiac(degree):
    signs = [(0, "Widder"), (30, "Stier"), (60, "Zwillinge"), (90, "Krebs"), (120, "Löwe"), (150, "Jungfrau"), (180, "Waage"), (210, "Skorpion"), (240, "Schütze"), (270, "Steinbock"), (300, "Wassermann"), (330, "Fische")]
    degree %= 360
    for i, (start, name) in enumerate(signs):
        if degree < start: return signs[i-1][1] if i > 0 else signs[0][1]
    return signs[-1][1]

# ===============================
# HAUPTLAUF
# ===============================
birth_pos = fetch_birth_positions(GEBURTSDATUM)
dates, year_data = fetch_year_positions(JAHR)

scores, mars_signs, venus_signs = [], [], []
for i in range(len(dates)):
    transit_pos = {p: year_data[p][i] for p in PLANET_IDS}
    scores.append(calculate_score(birth_pos, transit_pos))
    mars_signs.append(f"Mars in {get_zodiac(year_data['Mars'][i])}")
    venus_signs.append(f"Venus in {get_zodiac(year_data['Venus'][i])}")

df = pd.DataFrame({
    "Datum": [d.date() for d in dates],
    "Zodiac_Name": "",  # Spalte B für die 2. Achse
    "Score": scores,
    "Mars": mars_signs,
    "Venus": venus_signs
})

# Score normalisieren
s_min, s_max = df['Score'].min(), df['Score'].max()
df['Score_Prozent'] = ((df['Score'] - s_min) / (s_max - s_min)) * 100

# Vertikale Linien Logik & Zodiac Namen für Achse
df['Zodiac_Start'] = 0
for sign, start_md in ZODIAC_START.items():
    current_start_date = pd.to_datetime(f"{JAHR}-{start_md}").date()
    mask = df['Datum'] == current_start_date
    if mask.any():
        idx = df.index[mask][0]
        df.at[idx, 'Zodiac_Start'] = 100
        df.at[idx, 'Zodiac_Name'] = sign # Name wird in Spalte B eingetragen

# ===============================
# Excel mit mehrstufiger Achse
# ===============================
excel_path = f"Astrologie_Analyse_{JAHR}_DoppelAchse.xlsx"
df.to_excel(excel_path, index=False)

wb = load_workbook(excel_path)
ws = wb.active

chart = LineChart()
chart.title = f"Astrologischer Score {JAHR} (mit Tierkreis-Achse)"
chart.y_axis.title = "Score in %"

# Daten: Score (Spalte F) und Nadeln (Spalte G)
score_ref = Reference(ws, min_col=6, min_row=1, max_row=len(df)+1)
chart.add_data(score_ref, titles_from_data=True)

marker_ref = Reference(ws, min_col=7, min_row=1, max_row=len(df)+1)
chart.add_data(marker_ref, titles_from_data=True)

# --- MEHRSTUFIGE X-ACHSE ---
# Wir nehmen Spalte A (Datum) UND Spalte B (Zodiac_Name) zusammen!
categories = Reference(ws, min_col=1, max_col=2, min_row=2, max_row=len(df)+1)
chart.set_categories(categories)

# Optik der Nadeln
chart.series[1].graphicalProperties.line.solidFill = "FF0000"
chart.series[1].graphicalProperties.line.dashStyle = "sysDash"

ws.add_chart(chart, "I5")
wb.save(excel_path)

print(f"Fertig! In '{excel_path}' bilden Datum und Sternzeichen nun eine gemeinsame Achse.")